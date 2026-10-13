"""Read-only, source-located PDF/XLSX and installed JAR help queries."""
import argparse
import html
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile


def pdf_query(args):
    from pypdf import PdfReader
    reader = PdfReader(args.file)
    if not args.pages:
        return {'file': str(args.file.resolve()), 'pages': len(reader.pages),
                'metadata': {str(k): str(v) for k, v in (reader.metadata or {}).items()},
                'hint': 'Use --pages 1-3 or a bounded chapter range; PDF page numbers are 1-based'}
    chosen = set()
    for part in args.pages.split(','):
        limits = part.split('-')
        start, end = int(limits[0]), int(limits[-1])
        if len(limits) > 2 or start < 1 or end < start or end > len(reader.pages):
            raise ValueError('Invalid 1-based PDF page range')
        chosen.update(range(start, end + 1))
    if len(chosen) > 60:
        raise ValueError('Limit each query to 60 pages; use chapter index to narrow it')
    results = []
    for number in sorted(chosen):
        value = reader.pages[number - 1].extract_text() or ''
        if not args.find or args.find.casefold() in value.casefold():
            results.append({'pdf_page': number, 'text': value})
    return {'file': str(args.file.resolve()), 'results': results,
            'note': 'Text extraction may lose table alignment; inspect original page for ambiguous tables'}


def xlsx_query(args):
    import openpyxl
    from openpyxl.utils import get_column_letter
    # Normal mode is needed for authoritative merged ranges. Never save workbook.
    book = openpyxl.load_workbook(args.file, data_only=False, read_only=False)
    try:
        if not args.sheet:
            return {'file': str(args.file.resolve()), 'sheets': [
                {'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in book]}
        if args.sheet not in book.sheetnames:
            raise ValueError(f'Unknown sheet: {args.sheet}; list sheets first')
        sheet = book[args.sheet]
        merged = list(sheet.merged_cells.ranges)

        def cell_info(row, col):
            target = sheet.cell(row, col)
            source = target
            for area in merged:
                if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
                    source = sheet.cell(area.min_row, area.min_col)
                    break
            if source.value is None:
                return None
            result = {'cell': f'{get_column_letter(col)}{row}', 'value': source.value}
            if source.coordinate != target.coordinate:
                result['source_cell'] = source.coordinate
            if source.data_type == 'f':
                result['formula_not_evaluated'] = True
            if source.comment:
                result['comment'] = source.comment.text
            return result

        if args.start < 1 or (args.end is not None and args.end < args.start) or args.limit < 1:
            raise ValueError('Invalid row bounds or limit')
        rows = []
        matches = 0
        last = min(args.end if args.end is not None else sheet.max_row, sheet.max_row)
        for row in range(args.start, last + 1):
            cells = [v for col in range(1, sheet.max_column + 1) if (v := cell_info(row, col)) is not None]
            if cells and (not args.find or any(args.find.casefold() in str(v['value']).casefold() for v in cells)):
                matches += 1
                if len(rows) < args.limit:
                    rows.append({'row': row, 'cells': cells})
        return {'file': str(args.file.resolve()), 'sheet': sheet.title,
                'merged_ranges': [str(m) for m in merged], 'matching_rows': matches,
                'truncated': matches > len(rows), 'rows': rows,
                'note': 'Raw candidates; verify headers, package, device applicability and footnotes'}
    finally:
        book.close()


def help_query(args):
    with ZipFile(args.file) as jar:
        if not args.entry:
            return {'jar': str(args.file.resolve()), 'entries': [n for n in jar.namelist()
                    if n.endswith(('.html', '.htm')) and (args.find or '').casefold() in n.casefold()]}
        value = jar.read(args.entry).decode('utf-8', errors='replace')
        value = re.sub(r'<(script|style)\b[^>]*>.*?</\1>', '', value, flags=re.S | re.I)
        value = re.sub(r'<(?:br\s*/?|/p|/div|/tr|/li|/h[1-6])>', '\n', value, flags=re.I)
        value = html.unescape(re.sub(r'<[^>]+>', ' ', value))
        return {'jar': str(args.file.resolve()), 'entry': args.entry,
                'text': '\n'.join(' '.join(line.split()) for line in value.splitlines() if line.strip())}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('pdf')
    p.add_argument('file', type=Path)
    p.add_argument('--pages')
    p.add_argument('--find')
    p = sub.add_parser('xlsx')
    p.add_argument('file', type=Path)
    p.add_argument('--sheet')
    p.add_argument('--find')
    p.add_argument('--start', type=int, default=1)
    p.add_argument('--end', type=int)
    p.add_argument('--limit', type=int, default=30)
    p = sub.add_parser('help')
    p.add_argument('file', type=Path)
    p.add_argument('--entry')
    p.add_argument('--find')
    args = parser.parse_args()
    try:
        result = {'pdf': pdf_query, 'xlsx': xlsx_query, 'help': help_query}[args.command](args)
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        return 0
    except ImportError as exc:
        print(f'Missing dependency: {exc}. Use bundled Python with pypdf/openpyxl or an approved environment.', file=sys.stderr)
        return 2
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
    raise SystemExit(main())
